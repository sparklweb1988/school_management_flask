from flask import (
    Blueprint,
    render_template,
    request,
    redirect,
    url_for
)

from flask_login import (
    login_required,
    current_user
)

from extensions import db

from models import (
    FeePayment,
    Student,
    SchoolFee
)

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from flask import send_file

payment = Blueprint(
    "payment",
    __name__
)


# =========================================================
# PAYMENT LIST
# =========================================================

@payment.route("/fee-payments")
@login_required
def fee_payments():

    if current_user.role != "admin":
        return "Access denied", 403

    payments = FeePayment.query.order_by(
        FeePayment.payment_date.desc()
    ).all()

    return render_template(
        "feepayment/list.html",
        payments=payments
    )


# =========================================================
# CREATE PAYMENT
# =========================================================

@payment.route(
    "/fee-payments/create",
    methods=["GET", "POST"]
)
@login_required
def fee_payment_create():

    if current_user.role != "admin":
        return "Access denied", 403

    students = Student.query.order_by(
        Student.name
    ).all()

    fees = SchoolFee.query.order_by(
        SchoolFee.created_at.desc()
    ).all()

    if request.method == "POST":

        student_id = request.form["student_id"]

        fee_id = request.form["fee_id"]

        amount_paid = request.form["amount_paid"]

        reference = request.form["reference"]


        payment_record = FeePayment(

            student_id=student_id,

            fee_id=fee_id,

            amount_paid=amount_paid,

            reference=reference

        )


        db.session.add(payment_record)

        db.session.commit()


        return redirect(
            url_for("payment.fee_payments")
        )


    return render_template(
        "feepayment/create.html",

        students=students,

        fees=fees
    )


# =========================================================
# EDIT PAYMENT
# =========================================================

@payment.route(
    "/fee-payments/<int:id>/edit",
    methods=["GET", "POST"]
)
@login_required
def fee_payment_edit(id):

    if current_user.role != "admin":
        return "Access denied", 403


    payment_record = FeePayment.query.get_or_404(id)


    students = Student.query.order_by(
        Student.name
    ).all()

    fees = SchoolFee.query.order_by(
        SchoolFee.created_at.desc()
    ).all()


    if request.method == "POST":

        payment_record.student_id = request.form[
            "student_id"
        ]

        payment_record.fee_id = request.form[
            "fee_id"
        ]

        payment_record.amount_paid = request.form[
            "amount_paid"
        ]

        payment_record.reference = request.form[
            "reference"
        ]


        db.session.commit()


        return redirect(
            url_for("payment.fee_payments")
        )


    return render_template(
        "feepayment/edit.html",

        payment=payment_record,

        students=students,

        fees=fees
    )


# =========================================================
# DELETE PAYMENT
# =========================================================

@payment.route(
    "/fee-payments/<int:id>/delete"
)
@login_required
def fee_payment_delete(id):

    if current_user.role != "admin":
        return "Access denied", 403


    payment_record = FeePayment.query.get_or_404(id)


    db.session.delete(payment_record)

    db.session.commit()


    return redirect(
        url_for("payment.fee_payments")
    )






# =========================================================
# FEE PAYMENT REPORT
# =========================================================

@payment.route("/fee-payments/report")
@login_required
def fee_payment_report():

    if current_user.role != "admin":
        return "Access denied", 403

    payments = FeePayment.query.order_by(
        FeePayment.payment_date.desc()
    ).all()

    return render_template(
        "feepayment/report.html",
        payments=payments
    )





# =========================================================
# GENERATE FEE PAYMENT EXCEL REPORT
# =========================================================

@payment.route("/fee-payments/report/excel")
@login_required
def fee_payment_excel():

    if current_user.role != "admin":
        return "Access denied", 403


    payments = FeePayment.query.order_by(
        FeePayment.payment_date.desc()
    ).all()


    # -------------------------
    # CREATE EXCEL WORKBOOK
    # -------------------------

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Fee Payment Report"


    # -------------------------
    # REPORT TITLE
    # -------------------------

    worksheet["A1"] = "FEE PAYMENT REPORT"

    worksheet["A1"].font = Font(
        bold=True,
        size=16
    )

    worksheet["A1"].alignment = Alignment(
        horizontal="center"
    )


    worksheet.merge_cells(
        "A1:J1"
    )


    # -------------------------
    # HEADERS
    # -------------------------

    headers = [

        "Student",

        "Class",

        "Session",

        "Term",

        "School Fee",

        "Amount Paid",

        "Balance",

        "Status",

        "Reference",

        "Payment Date"

    ]


    for column, header in enumerate(
        headers,
        start=1
    ):

        cell = worksheet.cell(
            row=3,
            column=column
        )

        cell.value = header

        cell.font = Font(
            bold=True
        )


    # -------------------------
    # PAYMENT DATA
    # -------------------------

    row = 4


    for payment in payments:

        school_fee = payment.fee.amount

        amount_paid = payment.amount_paid

        balance = school_fee - amount_paid


        if balance <= 0:

            status = "Paid"

        else:

            status = "Balance"


        worksheet.cell(
            row=row,
            column=1
        ).value = payment.student.name


        worksheet.cell(
            row=row,
            column=2
        ).value = payment.fee.classroom.name


        worksheet.cell(
            row=row,
            column=3
        ).value = payment.fee.session.year


        worksheet.cell(
            row=row,
            column=4
        ).value = payment.fee.term.name


        worksheet.cell(
            row=row,
            column=5
        ).value = school_fee


        worksheet.cell(
            row=row,
            column=6
        ).value = amount_paid


        worksheet.cell(
            row=row,
            column=7
        ).value = balance


        worksheet.cell(
            row=row,
            column=8
        ).value = status


        worksheet.cell(
            row=row,
            column=9
        ).value = payment.reference


        if payment.payment_date:

            worksheet.cell(
                row=row,
                column=10
            ).value = payment.payment_date.strftime(
                "%d/%m/%Y %H:%M"
            )


        row += 1


    # -------------------------
    # COLUMN WIDTHS
    # -------------------------

    worksheet.column_dimensions["A"].width = 25
    worksheet.column_dimensions["B"].width = 20
    worksheet.column_dimensions["C"].width = 15
    worksheet.column_dimensions["D"].width = 15
    worksheet.column_dimensions["E"].width = 18
    worksheet.column_dimensions["F"].width = 18
    worksheet.column_dimensions["G"].width = 18
    worksheet.column_dimensions["H"].width = 15
    worksheet.column_dimensions["I"].width = 25
    worksheet.column_dimensions["J"].width = 22


    # -------------------------
    # SAVE TO MEMORY
    # -------------------------

    output = BytesIO()

    workbook.save(output)

    output.seek(0)


    # -------------------------
    # DOWNLOAD FILE
    # -------------------------

    return send_file(

        output,

        as_attachment=True,

        download_name="fee_payment_report.xlsx",

        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )

    )